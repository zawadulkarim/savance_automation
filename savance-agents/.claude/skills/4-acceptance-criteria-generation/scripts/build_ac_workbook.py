#!/usr/bin/env python3
"""
Generate an acceptance-criteria workbook in the house
"[Acceptance Criteria] ..." format: a 4-row identity header block, a section
title bar, a "Main Task:" (and optional "Sub Task:") banner, then the AC table

    ID | <group> | [Task] | Description | Acceptance Rules |
       Expected System Behavior | [Status] | Notes

with the group column merged over each contiguous block, exactly like the
four reference workbooks this format was learned from.

Two layouts, both observed in the reference files:

  per_ticket   (default) one "#N" sheet per ticket plus a Dashboard index --
               the "[Internal][Acceptance Criteria] Savance Q3 Release
               Items.xlsx" shape. Group column is "Module".
  single_sheet one named sheet covering many tickets, no Dashboard -- the
               "[Acceptance Criteria] Savance Miscellaneous (Phase 1).xlsx"
               shape. Usually paired with "task_column": true and
               "group_header": "Application".

Styling is deliberately the modern house palette shared with
build_checklist.py, NOT the reference files' raw Google-Sheets palette
(Calibri / #9CC2E5 / #3C78D8) -- see ../SKILL.md.

Usage:
    python build_ac_workbook.py <spec.json> <output.xlsx>

See ../SKILL.md for the JSON spec schema and the rules used to derive it.
"""
import sys
import json
import math

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FONT_NAME = "Segoe UI"

# Ink / chrome -- same constants as build_checklist.py so the acceptance
# criteria and the checklist read as one family of documents.
INK_PRIMARY = "0B0B0B"
INK_SECONDARY = "52514E"
BORDER_INK = "C3C2B7"
BAND_TINT = "F9F9F7"      # alternating group-block tint
GROUP_TINT = "E3EEFC"     # group-column tint (light accent-blue)
NEUTRAL_TINT = "F2F1EE"

ACCENT_DARK = "184F95"    # section title bars
ACCENT = "2A78D6"         # column headers
ACCENT_LIGHT = "CDE2FB"

REGULAR = Font(name=FONT_NAME, size=10, color=INK_PRIMARY)
BOLD = Font(name=FONT_NAME, size=10, bold=True, color=INK_PRIMARY)
BOLD_MUTED = Font(name=FONT_NAME, size=10, bold=True, color=INK_SECONDARY)
BOLD_GROUP = Font(name=FONT_NAME, size=10, bold=True, color=ACCENT_DARK)
HEADER_FILL = PatternFill("solid", fgColor=ACCENT)
HEADER_FONT = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
TITLE_FILL = PatternFill("solid", fgColor=ACCENT_DARK)
TITLE_FONT = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFF")
BANNER_FILL = PatternFill("solid", fgColor=ACCENT)
BANNER_FONT = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")

CENTER = Alignment(horizontal="center", vertical="center")
CENTER_WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)
TOP_WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)
GROUP_ALIGN = Alignment(horizontal="center", vertical="top", wrap_text=True)

THIN = Side(style="thin", color=BORDER_INK)
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

STATUS_COLORS = {
    "Passed": ("0CA30C", "FFFFFF"),
    "Failed": ("D03B3B", "FFFFFF"),
    "Untested": ("FAB219", "0B0B0B"),
}

# (key, header, width, alignment) -- letters are assigned left to right.
BASE_COLUMNS = [
    ("id", "ID", 10.0, "center"),
    ("group", "Module", 22.0, "group"),
    ("task", "Task", 30.0, "top"),
    ("description", "Description", 46.0, "top"),
    ("rules", "Acceptance Rules", 64.0, "top"),
    ("expected", "Expected System Behavior", 46.0, "top"),
    ("status", "Status", 12.0, "center"),
    ("notes", "Notes", 26.0, "top"),
]

ALIGNMENTS = {"center": CENTER_WRAP, "left": LEFT_WRAP,
              "top": TOP_WRAP, "group": GROUP_ALIGN}

DASHBOARD_WIDTHS = {"A": 9.0, "B": 54.0, "C": 18.0, "D": 16.0}
DASHBOARD_STATUSES = ["Not Started", "In Progress", "Completed", "On Hold"]

HEADER_BLOCK_ROWS = 4     # Project Name / Application Name / Prepared By / Last Updated
SECTION_TITLE_ROW = 6     # blank row 5 separates the block from the title bar


def column_plan(spec):
    """Ordered [(key, header, width, align, letter)] for this spec's shape."""
    keys = {"id", "group", "description", "rules", "expected", "notes"}
    if spec.get("task_column"):
        keys.add("task")
    if spec.get("status_column"):
        keys.add("status")
    group_header = spec.get("group_header", "Module")
    plan = []
    for key, header, width, align in BASE_COLUMNS:
        if key not in keys:
            continue
        if key == "group":
            header = group_header
        plan.append((key, header, width, align, get_column_letter(len(plan) + 1)))
    return plan


def est_lines(text, width):
    """Wrapped-line count for `text` in a span `width` chars wide.

    Only needed for **merged** cells: Excel auto-fits the height of a wrapped
    row it was given no height for, which is exact and therefore what the
    data rows rely on -- but it refuses to auto-fit a row containing a merged
    cell, so the merged header block would silently clip its longer values
    ("Savance Workplace Browser Interface, Kiosk, Visitor Management Outlook
    Addin" in the C:D span) without an explicit height.

    The estimate runs deliberately generous (0.9 chars per width unit): a
    too-tall row is merely airy, a too-short one hides text.
    """
    if text is None:
        return 1
    per = max(8.0, width * 0.9)
    lines = 0
    for segment in str(text).split("\n"):
        lines += max(1, math.ceil(len(segment) / per))
    return lines


def numbered(value):
    """Render a rules/expected value: a list becomes "1. ..\n2. ..", a string
    passes through so a single-sentence Expected System Behavior stays plain
    prose (as in the reference files).

    A one-item list still gets its "1. " prefix -- the reference files number
    single-rule cells too (e.g. AC2008's lone rule reads "1. This update of
    the workflow must be applied to both ..."), so dropping it for n=1 would
    read as an inconsistency across the sheet.
    """
    if value is None:
        return None
    if isinstance(value, str):
        return value
    items = [str(v).strip() for v in value if str(v).strip()]
    if not items:
        return None
    return "\n".join(f"{i}. {text}" for i, text in enumerate(items, 1))


def write_header_block(ws, plan, project_name, application_name, prepared_by,
                       last_updated, value_width):
    """Rows 1-4: label merged over the first two columns, value over the next
    two -- the reference files' identity block. `value_width` is the combined
    char width of the C:D value span, used to set an explicit height (merged
    cells are the one case Excel will not auto-fit)."""
    pairs = [("Project Name", project_name),
             ("Application Name", application_name),
             ("Prepared By", prepared_by),
             ("Last Updated", last_updated)]
    for i, (label_text, value) in enumerate(pairs, start=1):
        ws.merge_cells(f"A{i}:B{i}")
        cell = ws[f"A{i}"]
        cell.value = label_text
        cell.fill = TITLE_FILL
        cell.font = TITLE_FONT
        cell.alignment = CENTER_WRAP
        ws.merge_cells(f"C{i}:D{i}")
        value_cell = ws[f"C{i}"]
        value_cell.value = value
        value_cell.font = BOLD if i == 1 else REGULAR
        value_cell.alignment = LEFT_WRAP
        for col in ("A", "B", "C", "D"):
            ws[f"{col}{i}"].border = CELL_BORDER
        ws.row_dimensions[i].height = max(
            18.0, est_lines(value, value_width) * 14.5 + 4)
    return plan[-1][4]


def banner(ws, row, text, plan, fill, font):
    """A full-width merged bar (section title / Main Task / Sub Task)."""
    last_letter = plan[-1][4]
    ws.merge_cells(f"A{row}:{last_letter}{row}")
    cell = ws[f"A{row}"]
    cell.value = text
    cell.fill = fill
    cell.font = font
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    for _, _, _, _, letter in plan:
        ws[f"{letter}{row}"].border = CELL_BORDER
        ws[f"{letter}{row}"].fill = fill
    ws.row_dimensions[row].height = 20


def build_criteria_sheet(wb, sheet_name, spec, sheet_spec, index):
    """One AC sheet. `index` drives the AC ID prefix (AC{index}001...)."""
    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False
    plan = column_plan(spec)
    by_key = {key: letter for key, _, _, _, letter in plan}
    widths = {key: width for key, _, width, _, _ in plan}

    write_header_block(
        ws, plan,
        spec.get("project_name", "-"),
        sheet_spec.get("application_name", spec.get("application_name", "-")),
        spec.get("prepared_by", "Enosis QA"),
        spec.get("last_updated", "-"),
        value_width=plan[2][2] + plan[3][2],
    )

    section = sheet_spec.get("section_title") or spec.get("section_title")
    title_text = f"Acceptance Criteria: {section}" if section else "Acceptance Criteria"
    banner(ws, SECTION_TITLE_ROW, title_text, plan, TITLE_FILL, TITLE_FONT)

    row = SECTION_TITLE_ROW + 1
    main_task = sheet_spec.get("main_task")
    if main_task:
        banner(ws, row, f"Main Task: {main_task}", plan, BANNER_FILL, BANNER_FONT)
        row += 1
    sub_task = sheet_spec.get("sub_task")
    if sub_task:
        banner(ws, row, f"Sub Task: {sub_task}", plan, BANNER_FILL, BANNER_FONT)
        row += 1

    # --- table header ---
    header_row = row
    for _, header, _, _, letter in plan:
        cell = ws[f"{letter}{header_row}"]
        cell.value = header
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER_WRAP
        cell.border = CELL_BORDER
    ws.row_dimensions[header_row].height = 22

    # --- data rows ---
    first_row = header_row + 1
    r = first_row
    prev_group = None
    group_seq = 0
    group_start = r
    band_cols = [letter for key, _, _, _, letter in plan if key != "group"]

    def close_group(start, end):
        if end > start:
            ws.merge_cells(f"{by_key['group']}{start}:{by_key['group']}{end}")
        anchor = ws[f"{by_key['group']}{start}"]
        anchor.alignment = GROUP_ALIGN
        anchor.font = BOLD_GROUP

    criteria = sheet_spec.get("criteria", [])
    for i, item in enumerate(criteria, start=1):
        group_name = item.get(
            "module", item.get("application", item.get("group", "-")))
        if group_name != prev_group:
            group_seq += 1

        if group_seq % 2 == 0:
            band = PatternFill("solid", fgColor=BAND_TINT)
            for letter in band_cols:
                ws[f"{letter}{r}"].fill = band

        values = {
            "id": item.get("id") or f"AC{index}{i:03d}",
            "description": item.get("description"),
            "rules": numbered(item.get("rules")),
            "expected": numbered(item.get("expected")),
            "task": item.get("task"),
            "notes": item.get("notes"),
            "status": item.get("status", "Untested"),
        }
        for key, _, width, align, letter in plan:
            if key == "group":
                continue
            cell = ws[f"{letter}{r}"]
            cell.value = values.get(key)
            cell.font = REGULAR
            cell.alignment = ALIGNMENTS[align]

        # No explicit height: Excel auto-fits a wrapped row whose height was
        # never written, and auto-fit is exact where an estimate is not (see
        # est_lines' note). The reference workbooks carry no row heights
        # either.

        if group_name != prev_group:
            if prev_group is not None:
                close_group(group_start, r - 1)
            group_start = r
            ws[f"{by_key['group']}{r}"] = group_name
            prev_group = group_name

        r += 1

    last_row = r - 1
    if criteria:
        close_group(group_start, last_row)
        group_letter = by_key["group"]
        for row_ in range(first_row, r):
            ws[f"{group_letter}{row_}"].fill = PatternFill("solid", fgColor=GROUP_TINT)
        for row_ in range(header_row, r):
            for _, _, _, _, letter in plan:
                ws[f"{letter}{row_}"].border = CELL_BORDER

        if "status" in by_key:
            status_range = f"{by_key['status']}{first_row}:{by_key['status']}{last_row}"
            dv = DataValidation(type="list",
                                formula1='"Passed,Failed,Untested"',
                                allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(status_range)
            for value, (color, font_color) in STATUS_COLORS.items():
                # dxf fills read from start/end colour, not fgColor -- see
                # build_checklist.py's note; the fgColor spelling silently
                # renders as no fill at all in Excel.
                ws.conditional_formatting.add(
                    status_range,
                    CellIsRule(operator="equal", formula=[f'"{value}"'],
                               fill=PatternFill(start_color=color,
                                                end_color=color,
                                                fill_type="solid"),
                               font=Font(name=FONT_NAME, bold=True,
                                         color=font_color)))

    ws.freeze_panes = f"A{first_row}"
    for _, _, width, _, letter in plan:
        ws.column_dimensions[letter].width = width

    rules_total = sum(len(item.get("rules") or []) if not isinstance(
        item.get("rules"), str) else 1 for item in criteria)
    return {
        "sheet_name": sheet_name,
        "index": index,
        "main_task": main_task or sheet_spec.get("section_title") or sheet_name,
        "criteria": len(criteria),
        "rules": rules_total,
        "assigned": sheet_spec.get("assigned", "-"),
        "status": sheet_spec.get("status", "Not Started"),
        "url": sheet_spec.get("url"),
    }


def build_dashboard(wb, spec, metas):
    """Index sheet for the per_ticket layout: Item # | Asana Ticket |
    Assigned | Status, one row per ticket sheet."""
    ws = wb.create_sheet("Dashboard", 0)
    ws.sheet_view.showGridLines = False
    plan = [(letter.lower(), "", DASHBOARD_WIDTHS[letter], "center", letter)
            for letter in "ABCD"]
    write_header_block(
        ws, plan,
        spec.get("project_name", "-"),
        spec.get("application_name", "-"),
        spec.get("prepared_by", "Enosis QA"),
        spec.get("last_updated", "-"),
        value_width=DASHBOARD_WIDTHS["C"] + DASHBOARD_WIDTHS["D"],
    )

    header_row = SECTION_TITLE_ROW
    for letter, text in zip("ABCD", ["Item #", "Asana Ticket", "Assigned", "Status"]):
        cell = ws[f"{letter}{header_row}"]
        cell.value = text
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER_WRAP
    ws.row_dimensions[header_row].height = 22

    r = header_row + 1
    for i, meta in enumerate(metas, start=1):
        if i % 2 == 0:
            for letter in "ABCD":
                ws[f"{letter}{r}"].fill = PatternFill("solid", fgColor=BAND_TINT)
        for letter, value, align in (("A", f"{i}.0", CENTER),
                                     ("B", meta["main_task"], LEFT_WRAP),
                                     ("C", meta["assigned"], CENTER),
                                     ("D", meta["status"], CENTER)):
            cell = ws[f"{letter}{r}"]
            cell.value = value
            cell.font = REGULAR
            cell.alignment = align
        if meta.get("url"):
            ws[f"B{r}"].hyperlink = meta["url"]
            ws[f"B{r}"].font = Font(name=FONT_NAME, size=10, color="1155CC",
                                    underline="single")
        # No explicit height -- these cells are unmerged, so Excel auto-fits.
        r += 1

    if metas:
        dv = DataValidation(type="list",
                            formula1='"' + ",".join(DASHBOARD_STATUSES) + '"',
                            allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"D{header_row + 1}:D{r - 1}")
        for row_ in range(header_row, r):
            for letter in "ABCD":
                ws[f"{letter}{row_}"].border = CELL_BORDER

    ws.freeze_panes = f"A{header_row + 1}"
    for letter, width in DASHBOARD_WIDTHS.items():
        ws.column_dimensions[letter].width = width


def main():
    if len(sys.argv) != 3:
        print("Usage: python build_ac_workbook.py <spec.json> <output.xlsx>")
        sys.exit(1)

    spec_path, out_path = sys.argv[1], sys.argv[2]
    # utf-8-sig, not utf-8: anything written by PowerShell 5.1's Set-Content
    # -Encoding utf8 / Out-File carries a BOM, which plain utf-8 json.load
    # rejects outright ("Unexpected UTF-8 BOM").
    with open(spec_path, "r", encoding="utf-8-sig") as f:
        spec = json.load(f)

    layout = spec.get("layout", "per_ticket")
    if layout not in ("per_ticket", "single_sheet"):
        print(f"Unknown layout {layout!r}: use 'per_ticket' or 'single_sheet'")
        sys.exit(1)

    tickets = spec.get("tickets") or spec.get("sheets")
    if not tickets:
        print("Spec has no 'tickets' (per_ticket) / 'sheets' (single_sheet) entries")
        sys.exit(1)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    metas = []
    if layout == "per_ticket":
        start_index = spec.get("start_index", 1)
        for i, sheet_spec in enumerate(tickets):
            n = start_index + i
            metas.append(build_criteria_sheet(
                wb, sheet_spec.get("sheet_name", f"#{n}"), spec, sheet_spec, n))
        if spec.get("dashboard", True):
            build_dashboard(wb, spec, metas)
    else:
        for i, sheet_spec in enumerate(tickets, start=1):
            name = sheet_spec.get("sheet_name") or spec.get(
                "project_name", f"Sheet{i}")
            metas.append(build_criteria_sheet(wb, name[:31], spec, sheet_spec, i))

    wb.save(out_path)
    total_criteria = sum(m["criteria"] for m in metas)
    total_rules = sum(m["rules"] for m in metas)
    ratio = (total_rules / total_criteria) if total_criteria else 0
    print(f"Wrote {out_path}: {len(metas)} sheet(s)"
          f"{' + Dashboard' if layout == 'per_ticket' and spec.get('dashboard', True) else ''}"
          f" | {total_criteria} AC rows, {total_rules} acceptance rules "
          f"({ratio:.2f} rules/row)")
    for m in metas:
        print(f"  {m['sheet_name']}: {m['criteria']} AC rows, "
              f"{m['rules']} rules - {m['main_task'][:70]}")


if __name__ == "__main__":
    main()
